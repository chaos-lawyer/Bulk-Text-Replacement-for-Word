import os
from pathlib import Path
import shutil
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT / "src"))

from core.models import ExcelData, TemplateMergeItem
from core.template_merge import (
    build_template_context,
    create_merge_plan,
    generate_multi_template_batch,
    parse_output_directory_rule,
    plan_merge_jobs,
    render_output_directory,
    render_relative_folder,
)


class MultiTemplateMergeTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.work_dir = Path(self.temp_dir.name)

        # Create 2 Word templates
        self.tmpl1_path = self.work_dir / "采购合同.docx"
        doc1 = Document()
        doc1.add_paragraph("采购合同：甲方：{{甲方名称}}，合同编号：{{合同编号}}，金额：{{合同金额}}")
        doc1.save(self.tmpl1_path)

        self.tmpl2_path = self.work_dir / "保密协议.docx"
        doc2 = Document()
        doc2.add_paragraph("保密协议：甲方：{{甲方名称}}，保密期限：{{保密期限}}")
        doc2.save(self.tmpl2_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_parse_output_directory_rule(self):
        # 1. Path without variables
        r1 = parse_output_directory_rule("/Users/test/Output")
        self.assertEqual(r1.base_folder, os.path.normpath("/Users/test/Output"))
        self.assertEqual(r1.relative_rule, "")

        # 2. Path with variables
        r2 = parse_output_directory_rule("D:/Contracts/{{Region}}/{{Customer}}")
        self.assertEqual(r2.base_folder, os.path.normpath("D:/Contracts"))
        self.assertEqual(r2.relative_rule, "{{Region}}/{{Customer}}")

        # 3. Path with variable starting at root (invalid)
        r3 = parse_output_directory_rule("{{Region}}/Contracts")
        self.assertEqual(r3.base_folder, "")

    def test_render_output_directory_and_security(self):
        rule = parse_output_directory_rule(f"{self.work_dir}/{{{{地区}}}}/{{{{客户名称}}}}")
        values = {"地区": "华东", "客户名称": "示例科技"}
        target_dir, rel = render_output_directory(rule, values)

        self.assertEqual(rel, os.path.join("华东", "示例科技"))
        self.assertEqual(target_dir, (self.work_dir / "华东" / "示例科技").resolve())

        # Path traversal rejection
        values_evil = {"地区": "../../etc", "客户名称": "passwd"}
        target_evil, rel_evil = render_output_directory(rule, values_evil)
        # Guarantees destination stays under work_dir
        self.assertTrue(str(target_evil).startswith(str(self.work_dir.resolve())))

    def test_system_variables_in_template_context(self):
        tmpl = TemplateMergeItem(
            template_id="tmpl_1",
            file_path=str(self.tmpl1_path),
            display_name="采购合同模板",
            fields=["甲方名称", "合同编号", "合同金额"],
            filename_rule="{{模板名}}-{{合同编号}}",
            enabled=True,
        )
        row = {"甲方名称": "腾讯科技", "合同编号": "HT-001", "合同金额": "500万"}
        context = build_template_context(
            template=tmpl,
            row=row,
            excel_row=2,
            data_index=1,
            tmpl_index=1,
            today_date="2026-08-23",
        )

        self.assertEqual(context["模板名"], "采购合同模板")
        self.assertEqual(context["模板文件名"], "采购合同")
        self.assertEqual(context["模板所在文件夹"], str(self.tmpl1_path.parent))
        self.assertEqual(context["模板所在目录"], str(self.tmpl1_path.parent))
        self.assertEqual(context["模板序号"], "01")
        self.assertEqual(context["数据序号"], "001")
        self.assertEqual(context["Excel行号"], "2")
        self.assertEqual(context["日期"], "2026-08-23")
        self.assertEqual(context["甲方名称"], "腾讯科技")

    def test_plan_merge_jobs_per_template_rules(self):
        t1 = TemplateMergeItem(
            "t1", str(self.tmpl1_path), "采购合同",
            filename_rule="采购合同-{{编号}}", enabled=True
        )
        t2 = TemplateMergeItem(
            "t2", str(self.tmpl2_path), "保密协议",
            filename_rule="{{客户}}-保密协议-{{编号}}", enabled=True
        )
        data = ExcelData(
            headers=["客户", "编号", "金额", "期限"],
            rows=[
                {"客户": "公司A", "编号": "001", "金额": "10万", "期限": "3年"},
                {"客户": "公司B", "编号": "002", "金额": "20万", "期限": "5年"},
            ],
            excel_rows=[2, 3],
        )
        mapping = {
            "甲方名称": "客户",
            "合同编号": "编号",
            "合同金额": "金额",
            "保密期限": "期限",
        }

        jobs = plan_merge_jobs(
            templates=[t1, t2],
            data=data,
            mapping=mapping,
            output_directory_rule=f"{self.work_dir}/Output/{{{{客户}}}}",
        )

        # 2 templates × 2 rows = 4 jobs
        self.assertEqual(len(jobs), 4)

        # Check job 1: 公司A / 采购合同-001.docx
        self.assertEqual(jobs[0].relative_folder, "公司A")
        self.assertEqual(jobs[0].filename, "采购合同-001.docx")

        # Check job 2: 公司B / 采购合同-002.docx
        self.assertEqual(jobs[1].relative_folder, "公司B")
        self.assertEqual(jobs[1].filename, "采购合同-002.docx")

        # Check job 3: 公司A / 公司A-保密协议-001.docx
        self.assertEqual(jobs[2].relative_folder, "公司A")
        self.assertEqual(jobs[2].filename, "公司A-保密协议-001.docx")

    def test_generate_multi_template_batch_execution(self):
        t1 = TemplateMergeItem(
            "t1", str(self.tmpl1_path), "采购合同",
            filename_rule="采购合同-{{合同编号}}", enabled=True
        )
        t2 = TemplateMergeItem(
            "t2", str(self.tmpl2_path), "保密协议",
            filename_rule="保密协议-{{合同编号}}", enabled=True
        )
        data = ExcelData(
            headers=["甲方名称", "合同编号", "合同金额", "保密期限", "地区"],
            rows=[
                {"甲方名称": "华为技术", "合同编号": "HT-101", "合同金额": "100万", "保密期限": "2年", "地区": "华南"},
                {"甲方名称": "小米科技", "合同编号": "HT-102", "合同金额": "200万", "保密期限": "3年", "地区": "华北"},
            ],
            excel_rows=[2, 3],
        )
        mapping = {
            "甲方名称": "甲方名称",
            "合同编号": "合同编号",
            "合同金额": "合同金额",
            "保密期限": "保密期限",
        }
        out_folder = self.work_dir / "BatchOutput"

        results = generate_multi_template_batch(
            templates=[t1, t2],
            data=data,
            mapping=mapping,
            output_directory_rule=f"{out_folder}/{{{{地区}}}}/{{{{甲方名称}}}}",
        )

        self.assertEqual(len(results), 4)
        self.assertTrue(all(r.success for r in results))

        # Verify output files exist in subfolders
        file1 = out_folder / "华南" / "华为技术" / "采购合同-HT-101.docx"
        file2 = out_folder / "华北" / "小米科技" / "采购合同-HT-102.docx"
        file3 = out_folder / "华南" / "华为技术" / "保密协议-HT-101.docx"
        file4 = out_folder / "华北" / "小米科技" / "保密协议-HT-102.docx"

        self.assertTrue(file1.exists())
        self.assertTrue(file2.exists())
        self.assertTrue(file3.exists())
        self.assertTrue(file4.exists())

        # Verify content in generated file
        doc_huawei = Document(str(file1))
        self.assertIn("华为技术", doc_huawei.paragraphs[0].text)
        self.assertIn("100万", doc_huawei.paragraphs[0].text)

    def test_reused_plan_never_overwrites_file_created_after_preview(self):
        template = TemplateMergeItem(
            "t1",
            str(self.tmpl1_path),
            "采购合同",
            filename_rule=f"{self.work_dir}/Output/合同-{{{{合同编号}}}}",
            enabled=True,
        )
        data = ExcelData(
            headers=["合同编号"],
            rows=[{"合同编号": "001"}],
            excel_rows=[2],
        )
        mapping = {"合同编号": "合同编号"}
        plan = create_merge_plan([template], data, mapping)
        destination = Path(plan.jobs[0].destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b"external-file")

        results = generate_multi_template_batch(
            templates=[template],
            data=data,
            mapping=mapping,
            plan=plan,
        )

        self.assertEqual(len(results), 1)
        self.assertFalse(results[0].success)
        self.assertIn("预览后输出目标已存在", results[0].error)
        self.assertEqual(destination.read_bytes(), b"external-file")

    def test_execute_multi_template_merge_service_with_per_template_rules(self):
        """P0 Regression test: MergeService.execute_multi_template_merge succeeds when output_directory_rule is empty and templates have per-template rules."""
        from application.merge_service import MergeService

        out_dir_t1 = self.work_dir / "T1_Out"
        out_dir_t2 = self.work_dir / "T2_Out"

        t1 = TemplateMergeItem(
            "t1", str(self.tmpl1_path), "采购合同",
            filename_rule=f"{out_dir_t1}/采购-{{{{合同编号}}}}", enabled=True
        )
        t2 = TemplateMergeItem(
            "t2", str(self.tmpl2_path), "保密协议",
            filename_rule=f"{out_dir_t2}/保密-{{{{合同编号}}}}", enabled=True
        )
        data = ExcelData(
            headers=["甲方名称", "合同编号", "合同金额", "保密期限"],
            rows=[
                {"甲方名称": "华为技术", "合同编号": "HT-201", "合同金额": "100万", "保密期限": "2年"},
            ],
            excel_rows=[2],
        )
        mapping = {
            "甲方名称": "甲方名称",
            "合同编号": "合同编号",
            "合同金额": "合同金额",
            "保密期限": "保密期限",
        }

        res = MergeService.execute_multi_template_merge(
            templates=[t1, t2],
            excel_data=data,
            mapping=mapping,
            output_directory_rule="",
        )

        self.assertTrue(res.success, f"Expected success but got error: {res.error}")
        self.assertEqual(len(res.data), 2)
        self.assertTrue((out_dir_t1 / "采购-HT-201.docx").exists())
        self.assertTrue((out_dir_t2 / "保密-HT-201.docx").exists())

    def test_build_field_mapping_safety(self):
        """Test build_field_mapping only maps unique candidates and avoids ambiguous false positives."""
        from core.template_merge import build_field_mapping

        # Exact match
        res1 = build_field_mapping(["客户姓名", "合同金额"], ["客户姓名", "合同金额", "备注"])
        self.assertEqual(res1["客户姓名"], "客户姓名")
        self.assertEqual(res1["合同金额"], "合同金额")

        # Ambiguous substring: "名称" matches both "甲方名称" and "乙方名称" -> should NOT auto-map
        res2 = build_field_mapping(["名称"], ["甲方名称", "乙方名称", "备注"])
        self.assertEqual(res2["名称"], "")

        # Unique substring: "期限" uniquely matches "保密期限" -> maps
        res3 = build_field_mapping(["期限"], ["甲方名称", "保密期限", "备注"])
        self.assertEqual(res3["期限"], "保密期限")
