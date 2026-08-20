"""Create the checked-in Word/Excel fixtures used for manual and automated QA."""

from datetime import date
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "tests" / "samples"


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def create_template(path: Path):
    document = Document()
    title = document.add_paragraph("合同模板")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph("甲方：{{甲方名称}}；乙方：{{乙方名称}}")
    document.add_paragraph("Amount / 金额：{{合同金额}}")

    bold = document.add_paragraph("粗体字段：")
    bold_run = bold.add_run("{{签订日期}}")
    bold_run.bold = True

    split = document.add_paragraph("拆分字段：")
    first = split.add_run("{{管辖")
    first.font.name = "Arial"
    first.font.size = Pt(13)
    second = split.add_run("法院}}")
    second.font.name = "SimSun"
    second.italic = True

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "中文变量"
    table.cell(0, 1).text = "English variable"
    table.cell(1, 0).text = "{{甲方名称}}"
    table.cell(1, 1).text = "{{Contract ID}}"

    hyperlink_paragraph = document.add_paragraph("链接显示文字：")
    add_hyperlink(hyperlink_paragraph, "{{乙方名称}}", "https://example.com")

    section = document.sections[0]
    section.header.paragraphs[0].text = "页眉：{{Contract ID}}"
    section.footer.paragraphs[0].text = "页脚：{{签订日期}}"
    document.save(path)


def create_excel(path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合同数据"
    sheet.append(["甲方名称", "乙方名称", "合同金额", "签订日期", "管辖法院", "Contract ID", "多余列"])
    sheet.append(["A公司", "B公司", "100万元", date(2026, 8, 20), "太原市中级人民法院", "CN-001", "ignored"])
    sheet.append(["C公司", "D公司", 2000000, date(2026, 8, 21), "北京市第一中级人民法院", "CN-002", "ignored"])
    sheet.append(["E公司", "F公司", "", date(2026, 8, 22), "上海市第一中级人民法院", "CN-003", "ignored"])
    sheet.append([None, None, None, None, None, None, None])
    workbook.save(path)


def main():
    SAMPLES.mkdir(parents=True, exist_ok=True)
    create_template(SAMPLES / "合同模板.docx")
    create_excel(SAMPLES / "合同数据.xlsx")
    print(SAMPLES)


if __name__ == "__main__":
    sys.path.insert(0, str(ROOT / "src"))
    main()
