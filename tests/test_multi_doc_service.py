import os
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from application.multi_doc_service import MultiDocService
from application.task_models import CancellationToken, TaskState
from core.models import MultiDocItem


class MultiDocServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp_dir.name)

        self.doc_path = str(self.folder / "Test_Doc.docx")
        doc = Document()
        doc.add_paragraph("尊敬的{{客户姓名}}，您的订单{{订单号}}已发货。")
        doc.save(self.doc_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_validate_inputs_empty_and_invalid_paths(self):
        # Empty items
        err = MultiDocService.validate_inputs([])
        self.assertIsNotNone(err)
        self.assertIn("至少一个", err)

        # Non-existent file
        fake_item = MultiDocItem(file_path=str(self.folder / "ghost.docx"), filename="ghost.docx")
        err = MultiDocService.validate_inputs([fake_item])
        self.assertIsNotNone(err)
        self.assertIn("不存在", err)

        # In-place False with missing output folder
        valid_item = MultiDocItem(file_path=self.doc_path, filename="Test_Doc.docx")
        err = MultiDocService.validate_inputs([valid_item], in_place=False, output_folder="")
        self.assertIsNotNone(err)
        self.assertIn("输出保存目录", err)

    def test_scan_documents_and_cancellation(self):
        # Successful scan
        res = MultiDocService.scan_documents([self.doc_path])
        self.assertTrue(res.success)
        doc_vars, all_vars, doc_errors = res.data
        self.assertEqual(len(doc_errors), 0)
        self.assertEqual(all_vars, ["客户姓名", "订单号"])

        # Pre-cancelled token
        token = CancellationToken()
        token.cancel()
        res_cancelled = MultiDocService.scan_documents([self.doc_path], cancel_token=token)
        self.assertFalse(res_cancelled.success)
        self.assertEqual(res_cancelled.state, TaskState.CANCELLED)

    def test_preview_and_execute_batch_replace(self):
        items = [
            MultiDocItem(
                file_path=self.doc_path,
                filename="Test_Doc.docx",
                detected_variables=["客户姓名", "订单号"],
                replacements={"客户姓名": "张经理", "订单号": "ORD-20260822"},
            )
        ]

        # Preview
        preview_res = MultiDocService.generate_preview(items, ["客户姓名", "订单号"])
        self.assertTrue(preview_res.success)
        self.assertEqual(len(preview_res.data), 1)
        self.assertEqual(preview_res.data[0]["filename"], "Test_Doc.docx")

        # Execute in-place
        exec_res = MultiDocService.execute_batch_replace(items, in_place=True, create_backup=True)
        self.assertTrue(exec_res.success)
        self.assertEqual(exec_res.data.success_docs, 1)

        # Check replaced text
        doc_new = Document(self.doc_path)
        full_text = "".join(p.text for p in doc_new.paragraphs)
        self.assertIn("尊敬的张经理，您的订单ORD-20260822已发货。", full_text)


if __name__ == "__main__":
    unittest.main()
