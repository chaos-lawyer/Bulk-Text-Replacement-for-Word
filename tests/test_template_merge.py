import hashlib
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from template_merge import (  # noqa: E402
    ExcelData,
    build_field_mapping,
    build_output_filename,
    extract_template_fields,
    generate_batch,
    get_table_sheet_names,
    load_excel_data,
    load_table_data,
    mapped_row_values,
    normalize_excel_value,
    replace_docx_fields,
    sanitize_filename,
)


SAMPLES = ROOT / "tests" / "samples"


def document_text(path):
    document = Document(path)
    chunks = []
    for part in document.part.package.parts:
        if hasattr(part, "element"):
            for node in part.element.xpath(".//w:t"):
                chunks.append(node.text or "")
    return "".join(chunks)


class TemplateMergeTests(unittest.TestCase):
    def test_scan_excel_and_fields(self):
        fields = extract_template_fields(SAMPLES / "合同模板.docx")
        self.assertEqual(
            fields,
            ["甲方名称", "乙方名称", "合同金额", "签订日期", "管辖法院", "Contract ID"],
        )
        data = load_excel_data(SAMPLES / "合同数据.xlsx")
        self.assertEqual(len(data.rows), 3)
        self.assertEqual(data.rows[0]["签订日期"], "2026-08-20")
        self.assertEqual(data.rows[1]["合同金额"], "2000000")
        self.assertEqual(normalize_excel_value(None), "")

    def test_load_csv_utf8_and_gbk(self):
        with tempfile.TemporaryDirectory() as folder:
            # UTF-8 CSV
            csv_utf8 = Path(folder) / "data_utf8.csv"
            csv_utf8.write_text(
                "序号,甲方名称,乙方名称,合同金额\n1,甲公司,乙公司,5000\n2,丙公司,丁公司,8800\n",
                encoding="utf-8-sig",
            )
            data_u = load_table_data(csv_utf8)
            self.assertEqual(data_u.headers, ["序号", "甲方名称", "乙方名称", "合同金额"])
            self.assertEqual(len(data_u.rows), 2)
            self.assertEqual(data_u.rows[0]["甲方名称"], "甲公司")

            # GBK CSV
            csv_gbk = Path(folder) / "data_gbk.csv"
            csv_gbk.write_bytes(
                "序号,甲方名称,乙方名称,合同金额\n1,华北企业,华东企业,66000\n".encode("gbk")
            )
            data_g = load_table_data(csv_gbk)
            self.assertEqual(data_g.headers, ["序号", "甲方名称", "乙方名称", "合同金额"])
            self.assertEqual(len(data_g.rows), 1)
            self.assertEqual(data_g.rows[0]["甲方名称"], "华北企业")

    def test_batch_generate_with_csv_data(self):
        template = SAMPLES / "合同模板.docx"
        with tempfile.TemporaryDirectory() as folder:
            csv_path = Path(folder) / "data.csv"
            csv_path.write_text(
                "甲方名称,乙方名称,合同金额,签订日期,管辖法院,Contract ID\n"
                "测试甲,测试乙,500万,2026-08-21,北京市海淀区人民法院,CSV-001\n",
                encoding="utf-8",
            )
            data = load_table_data(csv_path)
            fields = extract_template_fields(template)
            mapping = build_field_mapping(fields, data.headers)
            results = generate_batch(
                template,
                data,
                mapping,
                folder,
                "{{Contract ID}}-{{甲方名称}}.docx",
            )
            self.assertEqual(len(results), 1)
            self.assertTrue(results[0].success)
            self.assertEqual(results[0].filename, "CSV-001-测试甲.docx")
            doc_content = document_text(Path(folder) / results[0].filename)
            self.assertIn("测试甲", doc_content)
            self.assertIn("CSV-001", doc_content)

    def test_standard_replace_preserves_split_run_format_and_hyperlink(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "output.docx"
            shutil.copy2(SAMPLES / "合同模板.docx", output)
            count = replace_docx_fields(
                output,
                {
                    "{{甲方名称}}": "A公司",
                    "{{乙方名称}}": "B公司",
                    "{{合同金额}}": "100万元",
                    "{{签订日期}}": "2026-08-20",
                    "{{管辖法院}}": "太原市中级人民法院",
                    "{{Contract ID}}": "CN-001",
                },
            )
            self.assertGreaterEqual(count, 9)
            self.assertNotIn("{{", document_text(output))
            document = Document(output)
            bold = next(p for p in document.paragraphs if p.text.startswith("粗体字段"))
            bold_replacement = next(run for run in bold.runs if "2026-08-20" in run.text)
            self.assertTrue(bold_replacement.bold)
            split = next(p for p in document.paragraphs if p.text.startswith("拆分字段"))
            replaced = next(run for run in split.runs if "太原市" in run.text)
            self.assertEqual(replaced.font.name, "Arial")
            self.assertEqual(replaced.font.size.pt, 13)
            rels = [rel for rel in document.part.rels.values() if rel.reltype.endswith("/hyperlink")]
            self.assertEqual(len(rels), 1)

    def test_batch_generates_three_files_without_changing_template(self):
        template = SAMPLES / "合同模板.docx"
        before = hashlib.sha256(template.read_bytes()).digest()
        data = load_excel_data(SAMPLES / "合同数据.xlsx")
        fields = extract_template_fields(template)
        mapping = build_field_mapping(fields, data.headers)
        with tempfile.TemporaryDirectory() as folder:
            results = generate_batch(
                template, data, mapping, folder,
                "{{甲方名称}}-{{乙方名称}}-合同.docx",
            )
            self.assertTrue(all(result.success for result in results))
            self.assertEqual(
                [result.filename for result in results],
                ["A公司-B公司-合同.docx", "C公司-D公司-合同.docx", "E公司-F公司-合同.docx"],
            )
            for result in results:
                self.assertNotIn("{{", document_text(Path(folder) / result.filename))
        self.assertEqual(before, hashlib.sha256(template.read_bytes()).digest())

    def test_empty_field_fallback_default_value(self):
        template = SAMPLES / "合同模板.docx"
        data = load_excel_data(SAMPLES / "合同数据.xlsx")
        fields = extract_template_fields(template)
        mapping = build_field_mapping(fields, data.headers)
        # Intentionally unmap "管辖法院" or set default
        defaults = {"管辖法院": "北京市海淀区人民法院", "Contract ID": "DEFAULT-999"}
        mapping["管辖法院"] = ""  # simulate unmapped / empty

        with tempfile.TemporaryDirectory() as folder:
            results = generate_batch(
                template,
                data,
                mapping,
                folder,
                "{{甲方名称}}-test.docx",
                default_values=defaults,
            )
            self.assertTrue(all(r.success for r in results))
            first_doc_text = document_text(Path(folder) / results[0].filename)
            self.assertIn("北京市海淀区人民法院", first_doc_text)
            self.assertNotIn("{{管辖法院}}", first_doc_text)

    def test_filename_rule_supports_excel_only_columns(self):
        """Verify that columns present ONLY in Excel (e.g. '序号', 'ID') can be used in filename rules."""
        row = {"序号": "001", "甲方名称": "测试甲方", "状态": "已审核"}
        mapping = {"甲方名称": "甲方名称"}  # '序号' and '状态' are not in Word template
        mapped = mapped_row_values(row, mapping)
        
        self.assertEqual(mapped.get("序号"), "001")
        self.assertEqual(mapped.get("状态"), "已审核")
        self.assertEqual(mapped.get("甲方名称"), "测试甲方")

        with tempfile.TemporaryDirectory() as folder:
            filename = build_output_filename("{{序号}}-{{甲方名称}}-{{状态}}.docx", mapped, folder)
            self.assertEqual(filename, "001-测试甲方-已审核.docx")

    def test_filename_safety_mapping_and_deduplication(self):
        self.assertEqual(sanitize_filename('A/B:合同?.docx'), "A_B_合同_.docx")
        with tempfile.TemporaryDirectory() as folder:
            reserved = set()
            first = build_output_filename("合同.docx", {}, folder, reserved=reserved)
            second = build_output_filename("合同.docx", {}, folder, reserved=reserved)
            self.assertEqual(first, "合同.docx")
            self.assertEqual(second, "合同 (2).docx")

    # ── Multiline paragraph-splitting tests ──────────────────────────────

    def _make_template(self, paragraphs_text: list[str], folder: str) -> str:
        """Create a minimal .docx template with the given paragraph texts."""
        doc = Document()
        # Clear default empty paragraph
        for p in doc.paragraphs:
            p._element.getparent().remove(p._element)
        for text in paragraphs_text:
            doc.add_paragraph(text)
        path = str(Path(folder) / "template.docx")
        doc.save(path)
        return path

    def test_multiline_standalone_field_splits_into_paragraphs(self):
        """A standalone {{field}} replaced with multi-line text produces multiple <w:p>."""
        with tempfile.TemporaryDirectory() as folder:
            path = self._make_template(["{{备注}}"], folder)
            output = Path(folder) / "output.docx"
            shutil.copy2(path, output)
            replace_docx_fields(output, {"{{备注}}": "第一行\n第二行\n第三行"})
            doc = Document(output)
            texts = [p.text for p in doc.paragraphs]
            self.assertEqual(texts, ["第一行", "第二行", "第三行"])

    def test_multiline_embedded_field_splits_paragraph(self):
        """A {{field}} embedded in other text still splits into paragraphs."""
        with tempfile.TemporaryDirectory() as folder:
            path = self._make_template(["甲方：{{甲方名称}} 签字"], folder)
            output = Path(folder) / "output.docx"
            shutil.copy2(path, output)
            replace_docx_fields(output, {"{{甲方名称}}": "张三\n李四"})
            doc = Document(output)
            texts = [p.text for p in doc.paragraphs]
            self.assertEqual(texts, ["甲方：张三", "李四 签字"])

    def test_multiline_inherits_run_format(self):
        """Split paragraphs inherit the character formatting of the placeholder run."""
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as folder:
            doc = Document()
            for p in doc.paragraphs:
                p._element.getparent().remove(p._element)
            p = doc.add_paragraph()
            run = p.add_run("{{内容}}")
            run.bold = True
            run.font.size = Pt(16)
            path = str(Path(folder) / "template.docx")
            doc.save(path)

            output = Path(folder) / "output.docx"
            shutil.copy2(path, output)
            replace_docx_fields(output, {"{{内容}}": "行A\n行B"})
            doc = Document(output)
            paras = doc.paragraphs
            self.assertEqual(len(paras), 2)
            self.assertEqual(paras[0].text, "行A")
            self.assertEqual(paras[1].text, "行B")
            # Both runs should inherit bold and font size
            self.assertTrue(paras[0].runs[0].bold)
            self.assertTrue(paras[1].runs[0].bold)
            self.assertEqual(paras[0].runs[0].font.size, Pt(16))
            self.assertEqual(paras[1].runs[0].font.size, Pt(16))

    def test_multiline_inherits_paragraph_alignment(self):
        """Split paragraphs inherit paragraph-level properties like alignment."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        with tempfile.TemporaryDirectory() as folder:
            doc = Document()
            for p in doc.paragraphs:
                p._element.getparent().remove(p._element)
            p = doc.add_paragraph("{{内容}}")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            path = str(Path(folder) / "template.docx")
            doc.save(path)

            output = Path(folder) / "output.docx"
            shutil.copy2(path, output)
            replace_docx_fields(output, {"{{内容}}": "居中行A\n居中行B\n居中行C"})
            doc = Document(output)
            paras = doc.paragraphs
            self.assertEqual(len(paras), 3)
            for p in paras:
                self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_no_newline_no_split(self):
        """Replacement without newlines does not split the paragraph."""
        with tempfile.TemporaryDirectory() as folder:
            path = self._make_template(["{{名称}}"], folder)
            output = Path(folder) / "output.docx"
            shutil.copy2(path, output)
            replace_docx_fields(output, {"{{名称}}": "普通文本"})
            doc = Document(output)
            self.assertEqual(len(doc.paragraphs), 1)
            self.assertEqual(doc.paragraphs[0].text, "普通文本")

    # ── replace_empty / keep_variable tests ──────────────────────────────

    def test_keep_variable_preserves_empty_field_placeholder(self):
        """With replace_empty=False, empty fields keep {{variable}} in output."""
        with tempfile.TemporaryDirectory() as folder:
            # Template has two fields, one mapped to empty value
            path = self._make_template(["{{甲方名称}}", "{{备注}}"], folder)
            data = ExcelData(
                headers=["甲方名称", "备注"],
                rows=[{"甲方名称": "A公司", "备注": ""}],
                excel_rows=[2],
            )
            mapping = {"甲方名称": "甲方名称", "备注": "备注"}
            results = generate_batch(
                path, data, mapping, folder, "test.docx", replace_empty=False,
            )
            self.assertTrue(results[0].success)
            doc = Document(Path(folder) / results[0].filename)
            texts = [p.text for p in doc.paragraphs]
            self.assertIn("A公司", texts)
            self.assertIn("{{备注}}", texts)

    def test_replace_empty_removes_placeholder(self):
        """With replace_empty=True (default), empty fields replace {{variable}} with empty."""
        with tempfile.TemporaryDirectory() as folder:
            path = self._make_template(["{{甲方名称}}", "{{备注}}"], folder)
            data = ExcelData(
                headers=["甲方名称", "备注"],
                rows=[{"甲方名称": "A公司", "备注": ""}],
                excel_rows=[2],
            )
            mapping = {"甲方名称": "甲方名称", "备注": "备注"}
            results = generate_batch(
                path, data, mapping, folder, "test.docx", replace_empty=True,
            )
            self.assertTrue(results[0].success)
            doc = Document(Path(folder) / results[0].filename)
            texts = [p.text for p in doc.paragraphs]
            self.assertIn("A公司", texts)
            self.assertNotIn("{{备注}}", texts)
            # Empty field should result in an empty paragraph
            self.assertIn("", texts)

    def test_keep_variable_with_default_still_uses_default(self):
        """With replace_empty=False, fields with defaults still use the default value."""
        with tempfile.TemporaryDirectory() as folder:
            path = self._make_template(["{{备注}}"], folder)
            data = ExcelData(
                headers=["备注"],
                rows=[{"备注": ""}],
                excel_rows=[2],
            )
            mapping = {"备注": "备注"}
            results = generate_batch(
                path, data, mapping, folder, "test.docx",
                replace_empty=False, default_values={"备注": "无"},
            )
            self.assertTrue(results[0].success)
            doc = Document(Path(folder) / results[0].filename)
            self.assertEqual(doc.paragraphs[0].text, "无")

    def test_per_field_empty_behaviors_can_be_mixed_in_one_document(self):
        """Each template variable can independently keep, clear, or use custom text."""
        with tempfile.TemporaryDirectory() as folder:
            path = self._make_template(
                ["{{保留字段}}", "{{清空字段}}", "{{自定义字段}}"], folder
            )
            data = ExcelData(
                headers=["保留字段", "清空字段", "自定义字段"],
                rows=[{"保留字段": "", "清空字段": "", "自定义字段": ""}],
                excel_rows=[2],
            )
            mapping = {field: field for field in data.headers}
            results = generate_batch(
                path,
                data,
                mapping,
                folder,
                "mixed.docx",
                default_values={"自定义字段": "待补充"},
                empty_field_behaviors={
                    "保留字段": "keep_variable",
                    "清空字段": "replace_empty",
                    "自定义字段": "custom",
                },
            )

            self.assertTrue(results[0].success)
            doc = Document(Path(folder) / results[0].filename)
            self.assertEqual(
                [paragraph.text for paragraph in doc.paragraphs],
                ["{{保留字段}}", "", "待补充"],
            )

    def test_multi_sheet_excel_loading(self):
        """Test creating an Excel workbook with multiple sheets and loading specific sheets."""
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            excel_path = Path(folder) / "multi_sheets.xlsx"
            wb = Workbook()
            
            # Sheet 1: 员工表
            ws1 = wb.active
            ws1.title = "员工表"
            ws1.append(["姓名", "部门", "职位"])
            ws1.append(["张三", "研发部", "工程师"])
            ws1.append(["李四", "市场部", "经理"])

            # Sheet 2: 客户表
            ws2 = wb.create_sheet(title="客户表")
            ws2.append(["客户名称", "联系人", "电话"])
            ws2.append(["甲公司", "王总", "13800000000"])

            wb.save(excel_path)
            wb.close()

            # Test get_table_sheet_names
            sheet_names = get_table_sheet_names(excel_path)
            self.assertEqual(sheet_names, ["员工表", "客户表"])

            # Test loading default active sheet (员工表)
            data_default = load_excel_data(excel_path)
            self.assertEqual(data_default.headers, ["姓名", "部门", "职位"])
            self.assertEqual(len(data_default.rows), 2)
            self.assertEqual(data_default.rows[0]["姓名"], "张三")

            # Test loading specific sheet (客户表)
            data_sheet2 = load_excel_data(excel_path, sheet_name="客户表")
            self.assertEqual(data_sheet2.headers, ["客户名称", "联系人", "电话"])
            self.assertEqual(len(data_sheet2.rows), 1)
            self.assertEqual(data_sheet2.rows[0]["客户名称"], "甲公司")

            # Test load_table_data with sheet_name
            data_table = load_table_data(excel_path, sheet_name="客户表")
            self.assertEqual(data_table.headers, ["客户名称", "联系人", "电话"])

    def test_get_table_sheet_names_single_and_csv(self):
        """Test get_table_sheet_names for single-sheet Excel and CSV files."""
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            # Single sheet Excel
            excel_path = Path(folder) / "single_sheet.xlsx"
            wb = Workbook()
            wb.save(excel_path)
            wb.close()
            self.assertEqual(get_table_sheet_names(excel_path), ["Sheet"])

            # CSV file
            csv_path = Path(folder) / "test.csv"
            csv_path.write_text("a,b\n1,2\n", encoding="utf-8")
            self.assertEqual(get_table_sheet_names(csv_path), [])

            # Non-existent file
            self.assertEqual(get_table_sheet_names(Path(folder) / "missing.xlsx"), [])


if __name__ == "__main__":
    unittest.main()
