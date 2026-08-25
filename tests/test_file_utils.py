import os
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(ROOT / "src"))

from core.file_utils import atomic_save_docx, create_safe_backup, verify_docx_integrity

SAMPLES = ROOT / "tests" / "samples"


class FileUtilsTests(unittest.TestCase):
    def test_create_safe_backup_non_overwriting(self):
        with tempfile.TemporaryDirectory() as folder:
            doc_file = Path(folder) / "contract.docx"
            doc_file.write_text("v1", encoding="utf-8")

            # 1st backup
            b1 = create_safe_backup(doc_file)
            self.assertEqual(Path(b1).name, "contract.docx.backup")
            self.assertTrue(Path(b1).exists())
            self.assertEqual(Path(b1).read_text(encoding="utf-8"), "v1")

            # Modify source and create 2nd backup
            doc_file.write_text("v2", encoding="utf-8")
            b2 = create_safe_backup(doc_file)
            self.assertEqual(Path(b2).name, "contract.docx.backup (2)")
            self.assertTrue(Path(b2).exists())
            self.assertEqual(Path(b2).read_text(encoding="utf-8"), "v2")

            # Verify 1st backup was preserved
            self.assertEqual(Path(b1).read_text(encoding="utf-8"), "v1")

            # 3rd backup
            doc_file.write_text("v3", encoding="utf-8")
            b3 = create_safe_backup(doc_file)
            self.assertEqual(Path(b3).name, "contract.docx.backup (3)")

    def test_atomic_save_docx_and_verification(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "output.docx"
            doc = Document()
            doc.add_paragraph("Hello World")

            atomic_save_docx(doc, target)
            self.assertTrue(target.exists())
            self.assertTrue(verify_docx_integrity(target))

            # Verify contents
            loaded = Document(target)
            self.assertEqual(loaded.paragraphs[0].text, "Hello World")
