import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

from docx import Document
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from core.models import ExcelData
from ui.main_window import MainWindow
from ui.pages.multi_doc_page import MultiDocPage

app = QApplication.instance() or QApplication(sys.argv)


class MultiDocGUITests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp_dir.name)

        # Create two sample Word files
        self.doc1 = str(self.folder / "Contract1.docx")
        self.doc2 = str(self.folder / "Contract2.docx")

        d1 = Document()
        d1.add_paragraph("甲方：{{甲方}}，乙方：{{乙方}}。")
        d1.save(self.doc1)

        d2 = Document()
        d2.add_paragraph("受聘人：{{受聘人}}，职务：{{职务}}。")
        d2.save(self.doc2)

        self.main_window = MainWindow()

    def tearDown(self):
        self.main_window.close()
        self.temp_dir.cleanup()

    def test_navigation_and_tab_switch(self):
        # Initial is replace page
        self.assertEqual(self.main_window.stacked_widget.currentWidget(), self.main_window.page_replace)

        # Switch to multi_doc
        self.main_window.nav.select_tab("multi_doc")
        self.assertEqual(self.main_window.stacked_widget.currentWidget(), self.main_window.page_multi_doc)

        # Switch to merge
        self.main_window.nav.select_tab("merge")
        self.assertEqual(self.main_window.stacked_widget.currentWidget(), self.main_window.page_merge)

    def test_multi_doc_page_add_documents_and_scan(self):
        page = self.main_window.page_multi_doc
        added = page.mapping_model.add_documents([self.doc1, self.doc2])
        self.assertEqual(added, 2)
        self.assertEqual(page.mapping_model.rowCount(), 2)

        # Scan variables synchronously in test by calling service directly or model
        doc_vars, all_vars = {
            self.doc1: ["甲方", "乙方"],
            self.doc2: ["受聘人", "职务"],
        }, ["甲方", "乙方", "受聘人", "职务"]
        page.mapping_model.set_variables(all_vars, doc_vars)

        self.assertEqual(page.mapping_model.columnCount(), 7)  # Index, Filename, 4 vars, Status

        # Set values in row 0
        page.mapping_model.setData(page.mapping_model.index(0, 2), "甲公司", Qt.EditRole)
        page.mapping_model.setData(page.mapping_model.index(0, 3), "乙公司", Qt.EditRole)

        # Check status of row 0
        status_0 = page.mapping_model.data(page.mapping_model.index(0, 6), Qt.DisplayRole)
        self.assertEqual(status_0, "✓ 数据就绪")

        # Check metrics
        ready, total = page.mapping_model.get_ready_metrics()
        self.assertEqual(ready, 1)
        self.assertEqual(total, 2)

    def test_multi_doc_excel_import_and_clear(self):
        page = self.main_window.page_multi_doc
        page.mapping_model.add_documents([self.doc1, self.doc2])
        page.mapping_model.set_variables(
            ["甲方", "乙方"],
            {self.doc1: ["甲方", "乙方"], self.doc2: ["甲方", "乙方"]},
        )

        excel_data = ExcelData(
            headers=["甲方", "乙方"],
            rows=[
                {"甲方": "甲1", "乙方": "乙1"},
                {"甲方": "甲2", "乙方": "乙2"},
            ],
            excel_rows=[2, 3],
        )
        page.excel_data = excel_data
        page._apply_excel_import(match_by_header=True)

        items = page.mapping_model.get_items()
        self.assertEqual(items[0].replacements["甲方"], "甲1")
        self.assertEqual(items[1].replacements["乙方"], "乙2")

        # Clear data
        page.mapping_model.clear_table_data()
        items_cleared = page.mapping_model.get_items()
        self.assertEqual(len(items_cleared[0].replacements), 0)
        self.assertEqual(page.mapping_model.rowCount(), 2)

    def test_multi_doc_multi_sheet_selection_behavior(self):
        """Test that multi-sheet Excel enables the sheet selector and single-sheet disables it."""
        from openpyxl import Workbook

        page = self.main_window.page_multi_doc

        # Single sheet Excel
        single_path = str(self.folder / "single.xlsx")
        wb1 = Workbook()
        wb1.active.title = "单表"
        wb1.save(single_path)
        wb1.close()

        page.edt_excel_path.setText(single_path)
        self.assertFalse(page.combo_sheet.isEnabled())
        self.assertFalse(page.lbl_sheet.isEnabled())
        self.assertEqual(page.combo_sheet.currentText(), "单表")

        # Multi sheet Excel
        multi_path = str(self.folder / "multi.xlsx")
        wb2 = Workbook()
        ws1 = wb2.active
        ws1.title = "SheetA"
        ws2 = wb2.create_sheet(title="SheetB")
        wb2.save(multi_path)
        wb2.close()

        page.edt_excel_path.setText(multi_path)
        self.assertTrue(page.combo_sheet.isEnabled())
        self.assertTrue(page.lbl_sheet.isEnabled())
        self.assertEqual(page.combo_sheet.count(), 2)
        self.assertEqual([page.combo_sheet.itemText(i) for i in range(2)], ["SheetA", "SheetB"])

    def test_merge_page_multi_sheet_selection_behavior(self):
        """Test that MergePage enables sheet selector on multi-sheet Excel and disables on single sheet/CSV."""
        from openpyxl import Workbook

        page = self.main_window.page_merge

        # CSV file
        csv_path = str(self.folder / "data.csv")
        Path(csv_path).write_text("col1,col2\n1,2\n", encoding="utf-8")
        page.edt_excel.setText(csv_path)
        self.assertFalse(page.combo_sheet.isEnabled())
        self.assertFalse(page.lbl_sheet.isEnabled())

        # Multi sheet Excel
        multi_path = str(self.folder / "multi_merge.xlsx")
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "第一批"
        ws2 = wb.create_sheet(title="第二批")
        wb.save(multi_path)
        wb.close()

        page.edt_excel.setText(multi_path)
        self.assertTrue(page.combo_sheet.isEnabled())
        self.assertTrue(page.lbl_sheet.isEnabled())
        self.assertEqual(page.combo_sheet.count(), 2)
        self.assertEqual([page.combo_sheet.itemText(i) for i in range(2)], ["第一批", "第二批"])

    def test_merge_page_scan_with_multi_sheet(self):
        """Verify scan_template on MergePage successfully processes multi-sheet Excel when clicked."""
        from openpyxl import Workbook

        page = self.main_window.page_merge

        tmpl_path = str(self.folder / "tmpl_scan.docx")
        doc = Document()
        doc.add_paragraph("甲方：{{甲方}}，乙方：{{乙方}}")
        doc.save(tmpl_path)

        multi_path = str(self.folder / "scan_sheets.xlsx")
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "无用表"
        ws1.append(["无关1", "无关2"])
        ws1.append(["val1", "val2"])

        ws2 = wb.create_sheet(title="合同数据")
        ws2.append(["甲方", "乙方"])
        ws2.append(["北京科技", "上海贸易"])
        wb.save(multi_path)
        wb.close()

        page.template_model.add_templates([tmpl_path])
        page.edt_excel.setText(multi_path)

        # Switch to second sheet
        page.combo_sheet.setCurrentIndex(1)
        self.assertEqual(page.combo_sheet.currentText(), "合同数据")

        # Trigger scan
        page.scan_templates()
        page._thread_pool.waitForDone(5000)
        app.processEvents()

        # Ensure scan succeeded and data is loaded
        self.assertIsNotNone(page.excel_data)
        self.assertEqual(page.excel_data.headers, ["甲方", "乙方"])
        self.assertEqual(len(page.excel_data.rows), 1)
        self.assertEqual(page.excel_data.rows[0]["甲方"], "北京科技")
        self.assertIn("合同数据", page.lbl_mapping_summary.text())


if __name__ == "__main__":
    unittest.main()
