import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from core.models import MultiDocItem
from core.multi_doc_replacer import (
    execute_multi_doc_preview,
    execute_multi_doc_replace,
    scan_documents_variables,
)

SAMPLES = ROOT / "tests" / "samples"


def get_doc_text(path: str | Path) -> str:
    doc = Document(path)
    chunks = []
    for part in doc.part.package.parts:
        if hasattr(part, "element"):
            for node in part.element.xpath(".//w:t"):
                chunks.append(node.text or "")
    return "".join(chunks)


class MultiDocReplacerTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp_dir.name)

        # Create 3 sample documents with distinct variables
        self.doc1_path = str(self.folder / "Doc1_Contract.docx")
        self.doc2_path = str(self.folder / "Doc2_NDA.docx")
        self.doc3_path = str(self.folder / "Doc3_Offer.docx")

        doc1 = Document()
        doc1.add_paragraph("甲方：{{甲方名称}}，乙方：{{乙方名称}}，金额：{{合同金额}}元。")
        doc1.save(self.doc1_path)

        doc2 = Document()
        doc2.add_paragraph("保密协议签署人：{{签署人}}，所属公司：{{甲方名称}}。")
        doc2.save(self.doc2_path)

        doc3 = Document()
        doc3.add_paragraph("录用通知书：{{签署人}} 先生/女士，入职岗位：{{岗位}}，薪资：{{月薪}}。")
        doc3.save(self.doc3_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_scan_documents_variables(self):
        file_paths = [self.doc1_path, self.doc2_path, self.doc3_path]
        doc_vars, all_vars, doc_errors = scan_documents_variables(file_paths)
        self.assertEqual(len(doc_errors), 0)

        self.assertEqual(doc_vars[self.doc1_path], ["甲方名称", "乙方名称", "合同金额"])
        self.assertEqual(doc_vars[self.doc2_path], ["签署人", "甲方名称"])
        self.assertEqual(doc_vars[self.doc3_path], ["签署人", "岗位", "月薪"])

        # Preserved order and deduplication
        self.assertEqual(all_vars, ["甲方名称", "乙方名称", "合同金额", "签署人", "岗位", "月薪"])

    def test_execute_multi_doc_preview(self):
        items = [
            MultiDocItem(
                file_path=self.doc1_path,
                filename="Doc1_Contract.docx",
                detected_variables=["甲方名称", "乙方名称", "合同金额"],
                replacements={"甲方名称": "腾讯科技", "乙方名称": "华为技术", "合同金额": "100万"},
            ),
            MultiDocItem(
                file_path=self.doc2_path,
                filename="Doc2_NDA.docx",
                detected_variables=["签署人", "甲方名称"],
                replacements={"签署人": "张三", "甲方名称": "腾讯科技"},
            ),
        ]
        all_vars = ["甲方名称", "乙方名称", "合同金额", "签署人"]
        preview = execute_multi_doc_preview(items, all_vars, output_folder=str(self.folder / "Out"))

        self.assertEqual(len(preview), 2)
        self.assertEqual(preview[0]["filename"], "Doc1_Contract.docx")
        self.assertEqual(len(preview[0]["fields"]), 4)
        self.assertTrue(preview[0]["fields"][0]["is_detected"])
        self.assertEqual(preview[0]["fields"][0]["value"], "腾讯科技")

    def test_execute_multi_doc_replace_to_output_dir(self):
        out_dir = str(self.folder / "Export")
        items = [
            MultiDocItem(
                file_path=self.doc1_path,
                filename="Doc1_Contract.docx",
                detected_variables=["甲方名称", "乙方名称", "合同金额"],
                replacements={"甲方名称": "阿里巴巴", "乙方名称": "蚂蚁金服", "合同金额": "500万"},
            ),
            MultiDocItem(
                file_path=self.doc2_path,
                filename="Doc2_NDA.docx",
                detected_variables=["签署人", "甲方名称"],
                replacements={"签署人": "李四", "甲方名称": "阿里巴巴"},
            ),
        ]

        res = execute_multi_doc_replace(items, output_folder=out_dir)

        self.assertEqual(res.total_docs, 2)
        self.assertEqual(res.success_docs, 2)
        self.assertEqual(res.failed_docs, 0)
        self.assertEqual(len(res.errors), 0)

        # Original files remain untouched
        doc1_orig_text = get_doc_text(self.doc1_path)
        self.assertIn("{{甲方名称}}", doc1_orig_text)

        # Output files contain replacements
        out_doc1 = str(Path(out_dir) / "Doc1_Contract.docx")
        out_doc2 = str(Path(out_dir) / "Doc2_NDA.docx")

        self.assertTrue(Path(out_doc1).exists())
        self.assertTrue(Path(out_doc2).exists())

        self.assertIn("甲方：阿里巴巴，乙方：蚂蚁金服，金额：500万元。", get_doc_text(out_doc1))
        self.assertIn("保密协议签署人：李四，所属公司：阿里巴巴。", get_doc_text(out_doc2))

    def test_execute_multi_doc_replace_in_place_with_backup(self):
        items = [
            MultiDocItem(
                file_path=self.doc3_path,
                filename="Doc3_Offer.docx",
                detected_variables=["签署人", "岗位", "月薪"],
                replacements={"签署人": "王五", "岗位": "架构师", "月薪": "50k"},
            )
        ]

        res = execute_multi_doc_replace(items, output_folder=None, create_backup=True)

        self.assertEqual(res.success_docs, 1)
        self.assertEqual(len(res.backup_files), 1)

        backup_file = self.doc3_path + ".backup"
        self.assertTrue(Path(backup_file).exists())
        self.assertIn("{{签署人}}", get_doc_text(backup_file))

        # Replaced content in place
        self.assertIn("录用通知书：王五 先生/女士，入职岗位：架构师，薪资：50k。", get_doc_text(self.doc3_path))

    def test_execute_multi_doc_replace_handles_missing_file_error(self):
        items = [
            MultiDocItem(
                file_path=str(self.folder / "Non_Existent.docx"),
                filename="Non_Existent.docx",
                replacements={"key": "val"},
            )
        ]

        res = execute_multi_doc_replace(items, output_folder=None)
        self.assertEqual(res.total_docs, 1)
        self.assertEqual(res.failed_docs, 1)
        self.assertEqual(len(res.errors), 1)
        self.assertIn("不存在", res.errors[0])


if __name__ == "__main__":
    unittest.main()
