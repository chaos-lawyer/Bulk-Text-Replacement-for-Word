"""Core services for Word template + Excel/CSV batch generation.

Deliberately decoupled from GUI frameworks for headless execution, automated testing, and thread safety.
Supports multi-template planning, system variables, subfolder rendering, and safe batch generation.
"""

from __future__ import annotations

import copy
import csv
import datetime as _dt
import os
from pathlib import Path
import re
import shutil
from typing import Callable, Iterable, Mapping, Optional

from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from openpyxl import load_workbook

from core.file_utils import atomic_save_docx
from core.models import ExcelData, MergeJob, MergePlanIssue, MergePlanResult, MergeResult, OutputDirectoryRule, TemplateMergeItem
from core.word_com import WordAutomationSession, traverse_story_ranges

FIELD_PATTERN = re.compile(r"\{\{([^{}\r\n]+)\}\}")
INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
INVALID_PATH_SEGMENT_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
INVISIBLE_CHARS_TABLE = str.maketrans("", "", "\u00ad\u200b\u200c\u200d\u2060\ufeff")
_WORD_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "application/vnd.ms-word.document.macroEnabled.main+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml",
}


def normalize_excel_value(value) -> str:
    """Convert an openpyxl / table value to predictable, user-facing text."""
    if value is None:
        return ""
    if isinstance(value, _dt.datetime):
        if value.time() == _dt.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, _dt.time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _load_csv_data(path: str | os.PathLike) -> ExcelData:
    """Read a CSV file with automatic encoding detection (utf-8-sig, utf-8, gb18030, gbk, latin1)."""
    encodings = ["utf-8-sig", "utf-8", "gb18030", "gbk", "latin1"]
    lines = None
    last_err = None

    for enc in encodings:
        try:
            with open(path, "r", encoding=enc, errors="strict") as f:
                lines = f.readlines()
            break
        except UnicodeDecodeError as err:
            last_err = err
            continue

    if lines is None:
        raise ValueError(f"无法以常见编码（UTF-8/GBK）解析 CSV 文件：{last_err}")

    reader = csv.reader(lines)
    rows_raw = list(reader)
    if not rows_raw:
        raise ValueError("CSV 文件为空。")

    first_row = rows_raw[0]
    headers = [normalize_excel_value(value).strip() for value in first_row]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or not any(headers):
        raise ValueError("CSV 第一行没有有效表头。")
    if any(not header for header in headers):
        raise ValueError("CSV 表头中间存在空白列。")
    duplicates = sorted({h for h in headers if headers.count(h) > 1})
    if duplicates:
        raise ValueError("CSV 表头重复：" + "、".join(duplicates))

    rows: list[dict[str, str]] = []
    excel_rows: list[int] = []
    for row_number, values in enumerate(rows_raw[1:], start=2):
        if len(values) < len(headers):
            values = values + [""] * (len(headers) - len(values))
        else:
            values = values[: len(headers)]

        normalized = [normalize_excel_value(value).strip() for value in values]
        if not any(value for value in normalized):
            continue
        rows.append(dict(zip(headers, normalized)))
        excel_rows.append(row_number)

    if not rows:
        raise ValueError("CSV 中没有有效数据行。")
    return ExcelData(headers, rows, excel_rows)


def get_table_sheet_names(path: str | os.PathLike) -> list[str]:
    """Retrieve sheet names from an Excel file, or [] for CSV / missing files."""
    if not os.path.exists(path):
        return []
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return []
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        try:
            wb = load_workbook(path, read_only=True, keep_links=False)
            try:
                return wb.sheetnames
            finally:
                wb.close()
        except Exception:
            return []
    return []


def load_table_data(path: str | os.PathLike, sheet_name: str | None = None) -> ExcelData:
    """Load table data from .xlsx, .xlsm, or .csv files into a unified ExcelData structure."""
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return _load_csv_data(path)
    elif suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return load_excel_data(path, sheet_name=sheet_name)
    else:
        raise ValueError(f"不支持的数据源格式：{suffix}。支持 .xlsx、.xlsm、.csv 格式。")


def load_excel_data(path: str | os.PathLike, sheet_name: str | None = None) -> ExcelData:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return _load_csv_data(path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active

        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            raise ValueError("Excel 工作表没有表头。")

        headers = [normalize_excel_value(value).strip() for value in first_row]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or not any(headers):
            raise ValueError("Excel 第一行没有有效表头。")
        if any(not header for header in headers):
            raise ValueError("Excel 表头中间存在空白列。")
        duplicates = sorted({h for h in headers if headers.count(h) > 1})
        if duplicates:
            raise ValueError("Excel 表头重复：" + "、".join(duplicates))

        rows: list[dict[str, str]] = []
        excel_rows: list[int] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True), start=2
        ):
            normalized = [normalize_excel_value(value) for value in values]
            if not any(value.strip() for value in normalized):
                continue
            rows.append(dict(zip(headers, normalized)))
            excel_rows.append(row_number)
        if not rows:
            raise ValueError("Excel 中没有有效数据行。")
        return ExcelData(headers, rows, excel_rows)
    finally:
        workbook.close()


def _iter_word_paragraphs(document) -> Iterable[Paragraph]:
    """Yield paragraphs in body, tables, headers, footers and text boxes once."""
    seen_parts: set[int] = set()
    for part in document.part.package.parts:
        if part.content_type not in _WORD_CONTENT_TYPES or not hasattr(part, "element"):
            continue
        if id(part) in seen_parts:
            continue
        seen_parts.add(id(part))
        parent = part.element
        for paragraph_element in parent.xpath(".//w:p"):
            yield Paragraph(paragraph_element, parent)


def _paragraph_runs(paragraph: Paragraph) -> list[Run]:
    """Include runs nested in hyperlinks, which Paragraph.runs omits."""
    return [Run(element, paragraph) for element in paragraph._p.xpath(".//w:r")]


def extract_template_fields(path: str | os.PathLike) -> list[str]:
    suffix = Path(path).suffix.lower()
    if suffix not in {".docx", ".docm"}:
        raise ValueError("快速模式扫描仅支持 .docx/.docm；.doc 请使用 Word COM 完整模式。")
    document = Document(path)
    fields: list[str] = []
    seen: set[str] = set()
    for paragraph in _iter_word_paragraphs(document):
        text = "".join(run.text for run in _paragraph_runs(paragraph)).translate(INVISIBLE_CHARS_TABLE)
        for match in FIELD_PATTERN.finditer(text):
            field = match.group(1)
            if field not in seen:
                seen.add(field)
                fields.append(field)
    return fields


def extract_template_fields_com(path: str | os.PathLike) -> list[str]:
    """Extract placeholder variable names from a Word document using Word COM."""
    fields: list[str] = []
    seen: set[str] = set()

    with WordAutomationSession() as word_app:
        doc = word_app.Documents.Open(os.path.abspath(path), ReadOnly=True)
        try:
            def _scan_story(story) -> int:
                text = story.Text
                if text:
                    for match in FIELD_PATTERN.finditer(text):
                        field = match.group(1)
                        if field not in seen:
                            seen.add(field)
                            fields.append(field)
                return 0

            traverse_story_ranges(doc, _scan_story)
        finally:
            doc.Close(0)

    return fields


def build_field_mapping(template_fields: list[str], excel_headers: list[str]) -> dict[str, str]:
    """Auto-map template variables to Excel/CSV headers with safe unique matching."""
    mapping: dict[str, str] = {}
    normalized_headers = {h.strip().lower(): h for h in excel_headers}

    for field in template_fields:
        field_clean = field.strip()
        field_lower = field_clean.lower()
        if field_clean in excel_headers:
            mapping[field] = field_clean
        elif field_lower in normalized_headers:
            mapping[field] = normalized_headers[field_lower]
        else:
            # Check unique substring candidate to avoid dangerous ambiguous mapping
            candidates = [h for h in excel_headers if field_clean in h.strip() or h.strip() in field_clean]
            if len(candidates) == 1:
                mapping[field] = candidates[0]
            else:
                mapping[field] = ""
    return mapping


def _replace_in_paragraph(paragraph: Paragraph, replacements: Mapping[str, str]) -> int:
    """Replace all placeholder fields in a paragraph while preserving formatting."""
    runs = _paragraph_runs(paragraph)
    if not runs:
        return 0

    full_text_chars = []
    char_map = []
    for run_idx, run in enumerate(runs):
        run_text = run.text or ""
        for offset, ch in enumerate(run_text):
            full_text_chars.append(ch)
            char_map.append((run_idx, offset))

    full_text = "".join(full_text_chars)
    if not full_text:
        return 0

    matches = []
    for field_placeholder, replacement_val in replacements.items():
        for m in re.finditer(re.escape(field_placeholder), full_text):
            matches.append((m.start(), m.end(), replacement_val))

    if not matches:
        return 0

    # Sort matches by start position in reverse order (right to left)
    matches.sort(key=lambda x: x[0], reverse=True)
    count = 0

    for start_char, end_char, replacement_str in matches:
        if start_char == end_char:
            continue
        start_run_idx, start_offset = char_map[start_char]
        end_run_idx, end_offset = char_map[end_char - 1]

        if start_run_idx == end_run_idx:
            target_run = runs[start_run_idx]
            orig_text = target_run.text
            target_run.text = orig_text[:start_offset] + replacement_str + orig_text[end_offset + 1:]
        else:
            first_run = runs[start_run_idx]
            first_run.text = first_run.text[:start_offset] + replacement_str
            for mid_idx in range(start_run_idx + 1, end_run_idx):
                runs[mid_idx].text = ""
            last_run = runs[end_run_idx]
            last_run.text = last_run.text[end_offset + 1:]
        count += 1

    return count


def _split_paragraph_at_newlines(paragraph: Paragraph) -> list[Paragraph]:
    """Split a paragraph into multiple paragraphs wherever runs contain newline characters."""
    has_newline = any("\n" in (r.text or "") or "\r" in (r.text or "") for r in paragraph.runs)
    if not has_newline:
        return [paragraph]

    p_elem = paragraph._p
    parent = p_elem.getparent()
    p_index = parent.index(p_elem)

    all_runs = paragraph.runs
    run_lines: list[list[tuple[Run, str]]] = [[]]

    for run in all_runs:
        text = run.text or ""
        lines = re.split(r"\r\n|\r|\n", text)
        for i, line in enumerate(lines):
            if i > 0:
                run_lines.append([])
            run_lines[-1].append((run, line))

    if len(run_lines) <= 1:
        return [paragraph]

    # First paragraph retains line 0
    for r in paragraph.runs:
        r.text = ""
    for orig_run, line_text in run_lines[0]:
        r_new = paragraph.add_run(line_text)
        _copy_run_format(orig_run, r_new)

    inserted = [paragraph]
    for seg_idx, segment in enumerate(run_lines[1:], start=1):
        new_p_elem = copy.deepcopy(p_elem)
        # Clear runs from cloned XML element
        for r_node in new_p_elem.xpath(".//w:r"):
            r_node.getparent().remove(r_node)

        parent.insert(p_index + seg_idx, new_p_elem)
        new_p = Paragraph(new_p_elem, paragraph._parent)
        for orig_run, line_text in segment:
            r_new = new_p.add_run(line_text)
            _copy_run_format(orig_run, r_new)
        inserted.append(new_p)

    return inserted


def _copy_run_format(src: Run, dst: Run) -> None:
    """Deep copy run properties (w:rPr) if available, preserving all font/color/style attributes."""
    try:
        from docx.oxml.ns import qn
        src_rPr = src._r.find(qn("w:rPr"))
        if src_rPr is not None:
            dst_rPr = dst._r.find(qn("w:rPr"))
            if dst_rPr is not None:
                dst._r.remove(dst_rPr)
            dst._r.append(copy.deepcopy(src_rPr))
            return
    except Exception:
        pass

    dst.bold = src.bold
    dst.italic = src.italic
    dst.underline = src.underline
    if src.font.name:
        dst.font.name = src.font.name
    if src.font.size:
        dst.font.size = src.font.size
    if src.font.color and src.font.color.rgb:
        dst.font.color.rgb = src.font.color.rgb


def replace_docx_fields(path: str | os.PathLike, replacements: Mapping[str, str]) -> int:
    document = Document(path)
    count = 0
    paragraphs = list(_iter_word_paragraphs(document))
    for paragraph in paragraphs:
        n = _replace_in_paragraph(paragraph, replacements)
        if n > 0:
            _split_paragraph_at_newlines(paragraph)
        count += n
    if count > 0:
        atomic_save_docx(document, path)
    return count


# ---------------- Path, Output Directory, and System Variable Helpers ----------------


def sanitize_path_segment(segment: str) -> str:
    """Sanitize a single directory or filename segment."""
    clean = INVALID_PATH_SEGMENT_CHARS.sub("_", segment).strip().rstrip(".")
    if not clean or clean in {".", ".."}:
        clean = "_"
    stem, suffix = os.path.splitext(clean)
    if stem.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        stem += "_"
    return stem[:180] + suffix


def sanitize_filename(filename: str, default_extension: str = ".docx") -> str:
    name = INVALID_FILENAME_CHARS.sub("_", filename).strip().rstrip(".")
    if not name:
        name = "Generated"

    stem, suffix = os.path.splitext(name)
    if not suffix or suffix.lower() != default_extension.lower():
        if suffix:
            stem = name[: -len(suffix)]
        name = stem + default_extension
        stem, suffix = os.path.splitext(name)

    if stem.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        stem += "_"
    return stem[:180] + suffix


def render_relative_folder(rule: str, values: Mapping[str, str]) -> str:
    """Render a sanitized relative subfolder path from a rule like '{{地区}}/{{客户名称}}'."""
    if not rule or not rule.strip():
        return ""

    rendered = FIELD_PATTERN.sub(lambda match: str(values.get(match.group(1), "")), rule)
    rendered = rendered.strip()
    if not rendered:
        return ""

    parts = re.split(r"[/\\]+", rendered)
    clean_parts = []
    for p in parts:
        p_clean = p.strip().rstrip(".")
        if not p_clean or p_clean in {".", ".."}:
            continue
        if ":" in p_clean:
            p_clean = p_clean.replace(":", "_")
        clean_seg = sanitize_path_segment(p_clean)
        if clean_seg and clean_seg != "_":
            clean_parts.append(clean_seg)

    return os.path.join(*clean_parts) if clean_parts else ""


def is_absolute_root(path_str: str) -> bool:
    """Check if a path string represents a valid platform-independent absolute root path."""
    if not path_str:
        return False
    norm = str(path_str).strip().replace("\\", "/")
    if norm.startswith("/") or norm.startswith("//"):
        return True
    if len(norm) >= 2 and norm[1] == ":" and norm[0].isalpha():
        return True
    return os.path.isabs(path_str)


def parse_output_directory_rule(rule: str | OutputDirectoryRule) -> OutputDirectoryRule:
    """Parse a single output directory rule string into a fixed base root and relative variable rule."""
    if isinstance(rule, OutputDirectoryRule):
        return rule

    raw = str(rule).strip()
    if not raw:
        return OutputDirectoryRule(raw_rule="", base_folder="", relative_rule="")

    normalized = raw.replace("\\", "/")
    parts = [p for p in normalized.split("/") if p]

    first_var_idx = -1
    for idx, part in enumerate(parts):
        if "{{" in part and "}}" in part:
            first_var_idx = idx
            break

    if first_var_idx == -1:
        base = os.path.normpath(raw)
        return OutputDirectoryRule(raw_rule=raw, base_folder=base, relative_rule="")

    if first_var_idx == 0:
        return OutputDirectoryRule(raw_rule=raw, base_folder="", relative_rule=raw)

    is_posix_abs = raw.startswith("/")
    base_parts = parts[:first_var_idx]
    if is_posix_abs:
        base_folder = "/" + "/".join(base_parts)
    elif ":" in base_parts[0]:
        base_folder = "/".join(base_parts)
    elif raw.startswith("\\\\") or raw.startswith("//"):
        base_folder = "//" + "/".join(base_parts)
    else:
        base_folder = "/".join(base_parts)

    base_folder = os.path.normpath(base_folder)
    rel_rule = "/".join(parts[first_var_idx:])
    return OutputDirectoryRule(raw_rule=raw, base_folder=base_folder, relative_rule=rel_rule)


def render_output_directory(
    dir_rule: str | OutputDirectoryRule, values: Mapping[str, str]
) -> tuple[Path, str]:
    """Calculate the safe destination directory from an OutputDirectoryRule. Returns (target_dir_path, relative_folder)."""
    parsed = parse_output_directory_rule(dir_rule)
    if not parsed.base_folder or not is_absolute_root(parsed.base_folder):
        raise ValueError(f"输出目录规则必须包含固定的绝对根目录：'{parsed.raw_rule}'")

    base_path = Path(parsed.base_folder).resolve()
    rel_folder = render_relative_folder(parsed.relative_rule, values)
    target_dir = (base_path / rel_folder).resolve()

    try:
        target_dir.relative_to(base_path)
    except ValueError:
        target_dir = base_path
        rel_folder = ""

    return target_dir, rel_folder


def build_template_context(
    template: TemplateMergeItem,
    row: Mapping[str, str],
    excel_row: int,
    data_index: int,
    tmpl_index: int = 1,
    mapping: Optional[Mapping[str, str]] = None,
    default_values: Optional[Mapping[str, str]] = None,
    empty_field_behaviors: Optional[Mapping[str, str]] = None,
    today_date: Optional[str] = None,
) -> dict[str, str]:
    """Construct replacement context dictionary containing Excel data, Word mapped variables, and system variables."""
    if today_date is None:
        today_date = _dt.date.today().isoformat()

    context = dict(row)

    if mapping:
        for field, header in mapping.items():
            raw_val = row.get(header, "") if header else ""
            should_rep, resolved_val, _ = resolve_empty_field(
                field=field,
                raw_value=raw_val,
                default_values=default_values,
                empty_field_behaviors=empty_field_behaviors,
                replace_empty=True,
            )
            if should_rep:
                context[field] = resolved_val
            else:
                context[field] = f"{{{{{field}}}}}"

    context["模板名"] = template.display_name
    context["模板文件名"] = Path(template.file_path).stem
    context["模板所在文件夹"] = str(Path(template.file_path).parent)
    context["模板所在目录"] = str(Path(template.file_path).parent)
    context["模板序号"] = f"{tmpl_index:02d}"
    context["数据序号"] = f"{data_index:03d}"
    context["Excel行号"] = str(excel_row)
    context["日期"] = today_date
    if "序号" not in context:
        context["序号"] = str(data_index)

    return context


def create_merge_plan(
    templates: list[TemplateMergeItem],
    data: ExcelData,
    mapping: Mapping[str, str],
    output_directory_rule: str | OutputDirectoryRule | os.PathLike = "",
    base_output_folder: str | os.PathLike | None = None,
    relative_folder_rule: str = "",
    default_filename_rule: str = "",
    default_values: Optional[Mapping[str, str]] = None,
    empty_field_behaviors: Optional[Mapping[str, str]] = None,
    check_existing: bool = True,
) -> MergePlanResult:
    """Plan an immutable MergePlanResult for all enabled templates and data rows."""
    if base_output_folder is not None:
        if relative_folder_rule:
            combined = f"{base_output_folder}/{relative_folder_rule}"
        else:
            combined = str(base_output_folder)
        dir_rule = parse_output_directory_rule(combined)
    else:
        dir_rule = parse_output_directory_rule(output_directory_rule)

    enabled_templates = [t for t in templates if t.enabled]
    if not enabled_templates:
        return MergePlanResult()

    jobs: list[MergeJob] = []
    issues: list[MergePlanIssue] = []
    batch_conflicts: list[dict] = []
    existing_file_conflicts: list[dict] = []
    reserved_paths: set[str] = set()
    today_str = _dt.date.today().isoformat()

    raw_jobs = []
    for tmpl_idx, tmpl in enumerate(enabled_templates, start=1):
        tmpl_fn_rule = (
            tmpl.filename_rule.strip()
            if tmpl.filename_rule and tmpl.filename_rule.strip()
            else (default_filename_rule.strip() or "{{模板名}}-{{数据序号}}")
        )
        ext = Path(tmpl.file_path).suffix.lower() or ".docx"

        tmpl_parent = str(Path(tmpl.file_path).parent)
        expanded_rule = (
            tmpl_fn_rule
            .replace("{{模板所在文件夹}}", tmpl_parent)
            .replace("{{模板所在目录}}", tmpl_parent)
            .replace("{{模板名}}", tmpl.display_name)
            .replace("{{模板文件名}}", Path(tmpl.file_path).stem)
        )

        normalized_rule = expanded_rule.replace("\\", "/")
        if "/" in normalized_rule:
            dir_part, fn_part = normalized_rule.rsplit("/", 1)
            tmpl_dir_rule = parse_output_directory_rule(dir_part)
            effective_fn_rule = fn_part
        else:
            tmpl_dir_rule = dir_rule
            effective_fn_rule = expanded_rule

        if not tmpl_dir_rule.base_folder:
            tmpl_dir_rule = parse_output_directory_rule(str(Path(tmpl.file_path).parent / "Generated"))

        for data_idx, (excel_row, row) in enumerate(zip(data.excel_rows, data.rows), start=1):
            context = build_template_context(
                template=tmpl,
                row=row,
                excel_row=excel_row,
                data_index=data_idx,
                tmpl_index=tmpl_idx,
                mapping=mapping,
                default_values=default_values,
                empty_field_behaviors=empty_field_behaviors,
                today_date=today_str,
            )

            target_dir, rel_folder = render_output_directory(tmpl_dir_rule, context)
            rendered_filename = FIELD_PATTERN.sub(
                lambda match: str(context.get(match.group(1), "")), effective_fn_rule
            )
            sanitized_fn = sanitize_filename(rendered_filename, ext)

            raw_jobs.append(
                {
                    "tmpl": tmpl,
                    "excel_row": excel_row,
                    "data_idx": data_idx,
                    "context": context,
                    "target_dir": target_dir,
                    "rel_folder": rel_folder,
                    "tmpl_fn_rule": tmpl_fn_rule,
                    "sanitized_fn": sanitized_fn,
                }
            )

    for rj in raw_jobs:
        target_dir = rj["target_dir"]
        sanitized_fn = rj["sanitized_fn"]
        stem, suffix = os.path.splitext(sanitized_fn)
        candidate_name = sanitized_fn
        number = 2
        collision_reason = ""
        conflict_kind = ""
        conflict_path = ""

        while True:
            dest_candidate = (target_dir / candidate_name).resolve()
            candidate_key = str(dest_candidate).casefold()
            exists_on_disk = check_existing and dest_candidate.exists()
            if candidate_key not in reserved_paths and not exists_on_disk:
                reserved_paths.add(candidate_key)
                final_dest = dest_candidate
                final_name = candidate_name
                break
            if candidate_key in reserved_paths:
                collision_reason = "同批任务重名"
                conflict_kind = "batch"
            else:
                collision_reason = "磁盘已有同名文件"
                conflict_kind = "existing"
            if not conflict_path:
                conflict_path = str(dest_candidate)
            candidate_name = f"{stem} ({number}){suffix}"
            number += 1

        if collision_reason:
            conflict_summary = {
                "requested": sanitized_fn,
                "assigned": final_name,
                "path": conflict_path,
            }
            if conflict_kind == "batch":
                batch_conflicts.append(conflict_summary)
            else:
                existing_file_conflicts.append(conflict_summary)
            issues.append(
                MergePlanIssue(
                    severity="warning",
                    code="NAME_COLLISION",
                    message=f"文件名冲突（{collision_reason}），已自动编号为：{final_name}",
                    template_id=rj["tmpl"].template_id,
                    template_name=rj["tmpl"].display_name,
                    excel_row=rj["excel_row"],
                    rule=rj["tmpl_fn_rule"],
                )
            )

        jobs.append(
            MergeJob(
                template_id=rj["tmpl"].template_id,
                template_path=rj["tmpl"].file_path,
                template_name=rj["tmpl"].display_name,
                excel_row=rj["excel_row"],
                data_index=rj["data_idx"],
                values=rj["context"],
                relative_folder=rj["rel_folder"],
                filename=final_name,
                destination=str(final_dest),
                filename_rule=rj["tmpl_fn_rule"],
                requested_filename=sanitized_fn,
                collision_reason=collision_reason,
            )
        )

    return MergePlanResult(
        jobs=jobs,
        issues=issues,
        batch_conflicts=batch_conflicts,
        existing_file_conflicts=existing_file_conflicts,
    )


def plan_merge_jobs(
    templates: list[TemplateMergeItem],
    data: ExcelData,
    mapping: Mapping[str, str],
    output_directory_rule: str | OutputDirectoryRule | os.PathLike = "",
    base_output_folder: str | os.PathLike | None = None,
    relative_folder_rule: str = "",
    default_filename_rule: str = "",
    default_values: Optional[Mapping[str, str]] = None,
    empty_field_behaviors: Optional[Mapping[str, str]] = None,
) -> list[MergeJob]:
    """Plan an immutable list of MergeJob instances for all enabled templates and data rows."""
    plan_res = create_merge_plan(
        templates=templates,
        data=data,
        mapping=mapping,
        output_directory_rule=output_directory_rule,
        base_output_folder=base_output_folder,
        relative_folder_rule=relative_folder_rule,
        default_filename_rule=default_filename_rule,
        default_values=default_values,
        empty_field_behaviors=empty_field_behaviors,
    )
    return plan_res.jobs


def mapped_row_values(
    row: Mapping[str, str],
    mapping: Mapping[str, str],
    default_values: Optional[Mapping[str, str]] = None,
    empty_field_behaviors: Optional[Mapping[str, str]] = None,
) -> dict[str, str]:
    """Expose template field names (including fallback defaults) to filename rules and replacements."""
    defaults = default_values or {}
    values = dict(row)
    for field, header in mapping.items():
        val = row.get(header, "") if header else ""
        should_replace, resolved, _ = resolve_empty_field(
            field, val, default_values=defaults, empty_field_behaviors=empty_field_behaviors
        )
        values[field] = resolved if should_replace else (f"{{{{{field}}}}}" if not val else val)
    return values


def build_output_filename(
    rule: str,
    row: Mapping[str, str],
    output_folder: str | os.PathLike,
    default_extension: str = ".docx",
    reserved: Optional[set[str]] = None,
) -> str:
    rendered = FIELD_PATTERN.sub(lambda match: str(row.get(match.group(1), "")), rule)
    filename = sanitize_filename(rendered, default_extension)
    reserved = reserved if reserved is not None else set()
    stem, suffix = os.path.splitext(filename)
    candidate = filename
    number = 2
    while candidate.casefold() in reserved or Path(output_folder, candidate).exists():
        candidate = f"{stem} ({number}){suffix}"
        number += 1
    reserved.add(candidate.casefold())
    return candidate


def resolve_empty_field(
    field: str,
    raw_value,
    default_values: Optional[Mapping[str, str]] = None,
    empty_field_behaviors: Optional[Mapping[str, str]] = None,
    replace_empty: bool = True,
) -> tuple[bool, str, str]:
    raw_text = str(raw_value) if raw_value is not None else ""
    if raw_text.strip():
        return True, raw_text, "value"

    defaults = default_values or {}
    behaviors = empty_field_behaviors or {}
    behavior = behaviors.get(field)

    if behavior == "custom":
        return True, str(defaults.get(field, "")), "custom"
    if behavior == "replace_empty":
        return True, "", "replace_empty"
    if behavior == "keep_variable":
        return False, "", "keep_variable"

    if field in defaults:
        return True, str(defaults[field]), "custom"
    return replace_empty, "", "replace_empty" if replace_empty else "keep_variable"


def _com_replace_document(document, replacements: Mapping[str, str]) -> int:
    def _replace_story(story) -> int:
        count = 0
        for search_text, replace_text in replacements.items():
            search_range = story.Duplicate
            search_range.Find.ClearFormatting()
            search_range.Find.Replacement.ClearFormatting()
            while search_range.Find.Execute(
                FindText=search_text,
                ReplaceWith=replace_text,
                Replace=1,
                Forward=True,
                Wrap=0,
                MatchCase=True,
                MatchWholeWord=False,
                MatchWildcards=False,
                MatchSoundsLike=False,
                MatchAllWordForms=False,
                Format=False,
            ):
                count += 1
        return count

    count, _warnings = traverse_story_ranges(document, _replace_story)
    return count


def generate_multi_template_batch(
    templates: list[TemplateMergeItem],
    data: ExcelData,
    mapping: Mapping[str, str],
    output_directory_rule: str | OutputDirectoryRule | os.PathLike = "",
    base_output_folder: str | os.PathLike | None = None,
    relative_folder_rule: str = "",
    default_filename_rule: str = "",
    use_com: bool = False,
    default_values: Optional[Mapping[str, str]] = None,
    empty_field_behaviors: Optional[Mapping[str, str]] = None,
    progress: Optional[Callable[[int, int, MergeResult], None]] = None,
    is_cancelled: Optional[Callable[[], bool]] = None,
    replace_empty: bool = True,
    plan: Optional[MergePlanResult] = None,
) -> list[MergeResult]:
    """Execute batch mail-merge across multiple templates and Excel/CSV data rows."""
    jobs = plan.jobs if plan is not None else plan_merge_jobs(
        templates=templates,
        data=data,
        mapping=mapping,
        output_directory_rule=output_directory_rule or base_output_folder or "",
        base_output_folder=base_output_folder,
        relative_folder_rule=relative_folder_rule,
        default_filename_rule=default_filename_rule,
        default_values=default_values,
        empty_field_behaviors=empty_field_behaviors,
    )

    if not jobs:
        return []

    # Ensure output directories exist for planned jobs
    unique_dirs = {Path(job.destination).parent for job in jobs}
    for d in unique_dirs:
        d.mkdir(parents=True, exist_ok=True)

    defaults = default_values or {}
    results: list[MergeResult] = []
    total_jobs = len(jobs)

    def _execute_jobs(word_app=None):
        for index, job in enumerate(jobs, start=1):
            if is_cancelled and is_cancelled():
                break

            result = MergeResult(
                excel_row=job.excel_row,
                filename=job.filename,
                success=False,
                template_id=job.template_id,
                template_name=job.template_name,
                output_path=job.destination,
                relative_folder=job.relative_folder,
            )

            copied_destination = False
            try:
                dest_path = Path(job.destination)
                dest_path.parent.mkdir(parents=True, exist_ok=True)
                if dest_path.exists():
                    raise FileExistsError(f"预览后输出目标已存在，为避免覆盖已停止写入：{dest_path}")

                shutil.copy2(job.template_path, str(dest_path))
                copied_destination = True

                replacements = {}
                for field, header in mapping.items():
                    val = job.values.get(field, "")
                    should_replace, resolved, _behavior = resolve_empty_field(
                        field,
                        val,
                        default_values=defaults,
                        empty_field_behaviors=empty_field_behaviors,
                        replace_empty=replace_empty,
                    )
                    if not should_replace:
                        continue
                    replacements[f"{{{{{field}}}}}"] = resolved

                if use_com:
                    doc = None
                    try:
                        doc = word_app.Documents.Open(os.path.abspath(str(dest_path)))
                        result.replacements = _com_replace_document(doc, replacements)
                        doc.Save()
                    finally:
                        if doc is not None:
                            doc.Close(False)
                else:
                    result.replacements = replace_docx_fields(str(dest_path), replacements)

                result.success = True
            except Exception as exc:
                result.error = str(exc)
                try:
                    if copied_destination and Path(job.destination).exists():
                        Path(job.destination).unlink()
                except OSError:
                    pass

            results.append(result)
            if progress:
                progress(index, total_jobs, result)

    if use_com:
        with WordAutomationSession() as word_app:
            _execute_jobs(word_app)
    else:
        _execute_jobs(None)

    return results


def generate_batch(
    template_path: str | os.PathLike,
    data: ExcelData,
    mapping: Mapping[str, str],
    output_folder: str | os.PathLike,
    filename_rule: str,
    use_com: bool = False,
    default_values: Optional[Mapping[str, str]] = None,
    progress: Optional[Callable[[int, int, MergeResult], None]] = None,
    is_cancelled: Optional[Callable[[], bool]] = None,
    replace_empty: bool = True,
    empty_field_behaviors: Optional[Mapping[str, str]] = None,
) -> list[MergeResult]:
    """Single template batch generation wrapper for backward compatibility."""
    tmpl_item = TemplateMergeItem(
        template_id="tmpl_1",
        file_path=str(template_path),
        display_name=Path(template_path).stem,
        filename_rule=filename_rule,
        enabled=True,
    )
    return generate_multi_template_batch(
        templates=[tmpl_item],
        data=data,
        mapping=mapping,
        output_directory_rule=output_folder,
        use_com=use_com,
        default_values=default_values,
        empty_field_behaviors=empty_field_behaviors,
        progress=progress,
        is_cancelled=is_cancelled,
        replace_empty=replace_empty,
    )
