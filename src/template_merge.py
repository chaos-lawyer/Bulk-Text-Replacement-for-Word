"""Core services for Word template + Excel batch generation.

The functions in this module deliberately do not depend on Tkinter.  This keeps
the merge workflow testable and lets the existing single-field GUI remain
unchanged apart from opening the new window.
"""

from __future__ import annotations

import datetime as _dt
import os
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, Mapping, Optional

from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from openpyxl import load_workbook


FIELD_PATTERN = re.compile(r"\{\{([^{}\r\n]+)\}\}")
INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
INVISIBLE_CHARS_TABLE = str.maketrans("", "", "\u00ad\u200b\u200c\u200d\u2060\ufeff")
_WORD_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "application/vnd.ms-word.document.macroEnabled.main+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml",
}


@dataclass
class ExcelData:
    headers: list[str]
    rows: list[dict[str, str]]
    excel_rows: list[int]


@dataclass
class MergeResult:
    excel_row: int
    filename: str
    success: bool
    replacements: int = 0
    error: str = ""


def normalize_excel_value(value) -> str:
    """Convert an openpyxl value to predictable, user-facing text."""
    if value is None:
        return ""
    if isinstance(value, _dt.datetime):
        if value.time() == _dt.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, _dt.time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def load_excel_data(path: str | os.PathLike) -> ExcelData:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            raise ValueError("Excel 文件没有表头。")

        headers = [normalize_excel_value(value).strip() for value in first_row]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or not any(headers):
            raise ValueError("Excel 第一行没有有效表头。")
        if any(not header for header in headers):
            raise ValueError("Excel 表头中间存在空白列。")
        duplicates = sorted({h for h in headers if headers.count(h) > 1})
        if duplicates:
            raise ValueError("Excel 表头重复：" + "、".join(duplicates))

        rows: list[dict[str, str]] = []
        excel_rows: list[int] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True), start=2
        ):
            normalized = [normalize_excel_value(value) for value in values]
            if not any(value.strip() for value in normalized):
                continue
            rows.append(dict(zip(headers, normalized)))
            excel_rows.append(row_number)
        if not rows:
            raise ValueError("Excel 中没有有效数据行。")
        return ExcelData(headers, rows, excel_rows)
    finally:
        workbook.close()


def _iter_word_paragraphs(document) -> Iterable[Paragraph]:
    """Yield paragraphs in body, tables, headers, footers and text boxes once."""
    seen_parts: set[int] = set()
    for part in document.part.package.parts:
        if part.content_type not in _WORD_CONTENT_TYPES or not hasattr(part, "element"):
            continue
        if id(part) in seen_parts:
            continue
        seen_parts.add(id(part))
        parent = part.element
        for paragraph_element in parent.xpath(".//w:p"):
            yield Paragraph(paragraph_element, parent)


def _paragraph_runs(paragraph: Paragraph) -> list[Run]:
    """Include runs nested in hyperlinks, which Paragraph.runs omits."""
    return [Run(element, paragraph) for element in paragraph._p.xpath(".//w:r")]


def extract_template_fields(path: str | os.PathLike) -> list[str]:
    suffix = Path(path).suffix.lower()
    if suffix not in {".docx", ".docm"}:
        raise ValueError("标准模式扫描仅支持 .docx/.docm；.doc 请使用 Word COM 高级模式。")
    document = Document(path)
    fields: list[str] = []
    seen: set[str] = set()
    for paragraph in _iter_word_paragraphs(document):
        text = "".join(run.text for run in _paragraph_runs(paragraph)).translate(INVISIBLE_CHARS_TABLE)
        for match in FIELD_PATTERN.finditer(text):
            field = match.group(1)
            if field not in seen:
                seen.add(field)
                fields.append(field)
    return fields


def extract_template_fields_com(path: str | os.PathLike) -> list[str]:
    """Scan legacy .doc or complex templates through an installed Word instance."""
    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("扫描 .doc 模板需要 Windows、Microsoft Word 和 pywin32。") from exc
    word_app = win32com.client.Dispatch("Word.Application")
    word_app.Visible = False
    word_app.DisplayAlerts = False
    document = None
    fields: list[str] = []
    seen: set[str] = set()
    try:
        document = word_app.Documents.Open(os.path.abspath(path), ReadOnly=True)
        for first_story in document.StoryRanges:
            story = first_story
            while story is not None:
                for match in FIELD_PATTERN.finditer(story.Text or ""):
                    field = match.group(1)
                    if field not in seen:
                        seen.add(field)
                        fields.append(field)
                try:
                    story = story.NextStoryRange
                except Exception:
                    story = None
        return fields
    finally:
        if document is not None:
            document.Close(False)
        word_app.Quit()


def build_field_mapping(fields: Iterable[str], headers: Iterable[str]) -> dict[str, str]:
    header_set = set(headers)
    return {field: field if field in header_set else "" for field in fields}


def _replace_in_paragraph(paragraph: Paragraph, replacements: Mapping[str, str]) -> int:
    runs = _paragraph_runs(paragraph)
    if not runs:
        return 0
    for run in runs:
        cleaned = run.text.translate(INVISIBLE_CHARS_TABLE)
        if cleaned != run.text:
            run.text = cleaned

    run_texts = [run.text for run in runs]
    text = "".join(run_texts)
    if not text:
        return 0
    keys = [key for key in replacements if key]
    if not keys:
        return 0
    pattern = re.compile("|".join(re.escape(key) for key in sorted(keys, key=len, reverse=True)))
    matches = list(pattern.finditer(text))
    if not matches:
        return 0

    starts: list[int] = []
    offset = 0
    for value in run_texts:
        starts.append(offset)
        offset += len(value)

    def locate(position: int, prefer_previous: bool = False) -> tuple[int, int]:
        for index, start in enumerate(starts):
            end = start + len(run_texts[index])
            if start <= position < end or (prefer_previous and position == end and end > start):
                return index, position - start
        return len(runs) - 1, len(runs[-1].text)

    # Work backwards so offsets before each match remain valid.  Replacement text
    # inherits the formatting of the first run occupied by the placeholder.
    for match in reversed(matches):
        start_index, start_offset = locate(match.start())
        end_index, end_offset = locate(match.end(), prefer_previous=True)
        value = str(replacements[match.group(0)])
        if start_index == end_index:
            current = runs[start_index].text
            runs[start_index].text = current[:start_offset] + value + current[end_offset:]
        else:
            prefix = runs[start_index].text[:start_offset]
            suffix = runs[end_index].text[end_offset:]
            runs[start_index].text = prefix + value
            for index in range(start_index + 1, end_index):
                runs[index].text = ""
            runs[end_index].text = suffix
    return len(matches)


def replace_docx_fields(path: str | os.PathLike, replacements: Mapping[str, str]) -> int:
    # python-docx keeps unknown related package parts when round-tripping; unlike
    # openpyxl its Document factory has no keep_vba argument.
    document = Document(path)
    count = 0
    for paragraph in _iter_word_paragraphs(document):
        count += _replace_in_paragraph(paragraph, replacements)
    document.save(path)
    return count


def sanitize_filename(filename: str, default_extension: str = ".docx") -> str:
    name = INVALID_FILENAME_CHARS.sub("_", filename).strip().rstrip(".")
    if not name:
        name = "Generated"
    if not Path(name).suffix:
        name += default_extension
    stem, suffix = os.path.splitext(name)
    # Windows also rejects these device names, even with an extension.
    if stem.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        stem += "_"
    return stem[:180] + suffix


def build_output_filename(
    rule: str, row: Mapping[str, str], output_folder: str | os.PathLike,
    default_extension: str = ".docx", reserved: Optional[set[str]] = None,
) -> str:
    rendered = FIELD_PATTERN.sub(lambda match: row.get(match.group(1), ""), rule)
    filename = sanitize_filename(rendered, default_extension)
    reserved = reserved if reserved is not None else set()
    stem, suffix = os.path.splitext(filename)
    candidate = filename
    number = 2
    while candidate.casefold() in reserved or Path(output_folder, candidate).exists():
        candidate = f"{stem} ({number}){suffix}"
        number += 1
    reserved.add(candidate.casefold())
    return candidate


def mapped_row_values(row: Mapping[str, str], mapping: Mapping[str, str]) -> dict[str, str]:
    """Expose template field names (including manually mapped ones) to filename rules."""
    values = dict(row)
    values.update({field: row.get(header, "") for field, header in mapping.items() if header})
    return values


def _com_replace_document(document, replacements: Mapping[str, str]) -> int:
    """Replace all fields across Word story ranges, preserving Word-managed content."""
    count = 0
    for first_story in document.StoryRanges:
        story = first_story
        while story is not None:
            for search_text, replace_text in replacements.items():
                search_range = story.Duplicate
                search_range.Find.ClearFormatting()
                search_range.Find.Replacement.ClearFormatting()
                while search_range.Find.Execute(
                    FindText=search_text, ReplaceWith=replace_text, Replace=1,
                    Forward=True, Wrap=0, MatchCase=True, MatchWholeWord=False,
                    MatchWildcards=False, MatchSoundsLike=False,
                    MatchAllWordForms=False, Format=False,
                ):
                    count += 1
            try:
                story = story.NextStoryRange
            except Exception:
                story = None
    return count


def generate_batch(
    template_path: str | os.PathLike,
    data: ExcelData,
    mapping: Mapping[str, str],
    output_folder: str | os.PathLike,
    filename_rule: str,
    use_com: bool = False,
    replace_empty: bool = True,
    progress: Optional[Callable[[int, int, MergeResult], None]] = None,
) -> list[MergeResult]:
    template_path = str(Path(template_path).resolve())
    output_folder = str(Path(output_folder).resolve())
    Path(output_folder).mkdir(parents=True, exist_ok=True)
    extension = Path(template_path).suffix.lower()
    if extension not in {".doc", ".docx", ".docm"}:
        raise ValueError("模板必须是 .doc、.docx 或 .docm 文件。")
    if not use_com and extension == ".doc":
        raise ValueError(".doc 模板必须使用 Word COM 高级模式。")

    word_app = None
    if use_com:
        try:
            import win32com.client
        except ImportError as exc:
            raise RuntimeError("高级模式需要 Windows、Microsoft Word 和 pywin32。") from exc
        word_app = win32com.client.Dispatch("Word.Application")
        word_app.Visible = False
        word_app.DisplayAlerts = False
        word_app.ScreenUpdating = False

    results: list[MergeResult] = []
    reserved: set[str] = set()
    try:
        total = len(data.rows)
        for index, (excel_row, row) in enumerate(zip(data.excel_rows, data.rows), start=1):
            filename = build_output_filename(
                filename_rule, mapped_row_values(row, mapping), output_folder, extension, reserved
            )
            destination = str(Path(output_folder, filename))
            result = MergeResult(excel_row, filename, False)
            try:
                shutil.copy2(template_path, destination)
                replacements = {}
                for field, header in mapping.items():
                    if not header:
                        continue
                    value = row.get(header, "")
                    if value or replace_empty:
                        replacements[f"{{{{{field}}}}}"] = value
                if use_com:
                    document = None
                    try:
                        document = word_app.Documents.Open(os.path.abspath(destination))
                        result.replacements = _com_replace_document(document, replacements)
                        document.Save()
                    finally:
                        if document is not None:
                            document.Close(False)
                else:
                    result.replacements = replace_docx_fields(destination, replacements)
                result.success = True
            except Exception as exc:
                result.error = str(exc)
                try:
                    if Path(destination).exists():
                        Path(destination).unlink()
                except OSError:
                    pass
            results.append(result)
            if progress:
                progress(index, total, result)
    finally:
        if word_app is not None:
            try:
                word_app.ScreenUpdating = True
                word_app.Quit()
            except Exception:
                pass
    return results
